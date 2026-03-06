import tkinter as tk
from tkinter import filedialog, messagebox
from selenium import webdriver
from selenium.webdriver.common.by import By
from selenium.webdriver.support.ui import WebDriverWait
from selenium.webdriver.support import expected_conditions as EC
from selenium.webdriver.chrome.service import Service
from selenium.webdriver.chrome.options import Options
from webdriver_manager.chrome import ChromeDriverManager
from openpyxl import load_workbook
from screeninfo import get_monitors  # pip install screeninfo
import threading
import time
# import os

# ================= 사용자 설정 =================
URL = "https://coupon.withhive.com/soulstrike?t=1737278324092"
# CHROMEDRIVER_PATH = r"C:\chromedriver\chromedriver.exe"  # 직접 다운로드한 경로

CS_CODE_INPUT_SELECTOR = "input#cs_code"
COUPON_INPUT_SELECTOR = "input#coupon_code"
SUBMIT_BUTTON_SELECTOR = "button.btn_use"
SLEEP_SEC = 1.5
# =================================================

def log_append(msg):
    """GUI 로그창에 메시지 추가"""
    txt_log.insert(tk.END, msg + "\n")
    txt_log.see(tk.END)  # 자동 스크롤
    root.update_idletasks()

def run_coupon_process(excel_path, coupon_code):
    try:
        wb = load_workbook(excel_path)
        ws = wb.active
        ids = [row[0].value for row in ws.iter_rows(min_row=2) if row[0].value]

        options = Options()
        # options.add_argument("--headless")  # 필요시 주석 해제
        options.add_argument("--no-sandbox")
        options.add_argument("--disable-dev-shm-usage")
        options.add_argument("--disable-gpu")

        log_append("✅ Chrome 실행 중...")
        # driver = webdriver.Chrome(service=Service(CHROMEDRIVER_PATH), options=options)
        driver = webdriver.Chrome(service=Service(ChromeDriverManager().install()), options=options)
        
        # 현재 모니터 해상도 가져오기
        monitor = get_monitors()[0]
        screen_width = monitor.width
        screen_height = monitor.height

        # 원하는 크기
        win_width = 100
        win_height = 100

        # 오른쪽 하단 위치 계산
        pos_x = screen_width - win_width - 20   # 오른쪽 여백 20px
        pos_y = screen_height - win_height - 80 # 아래 여백 (작업표시줄 고려해서 80px)

        # 창 위치/크기 설정
        driver.set_window_size(win_width, win_height)
        driver.set_window_position(pos_x, pos_y)
        # driver.set_window_position(200, 100)
        # driver.set_window_size(100, 100)
        driver.get(URL)
        wait = WebDriverWait(driver, 10)
        log_append("✅ 페이지 접속 완료.")

        for idx, cs_code in enumerate(ids, start=2):
            try:
                lbl_status.config(text=f"현재 처리 중: {cs_code}")
                log_append(f"[{idx-1}] {cs_code} 처리 중...")

                cs_input = wait.until(EC.presence_of_element_located((By.CSS_SELECTOR, CS_CODE_INPUT_SELECTOR)))
                cs_input.clear()
                cs_input.send_keys(str(cs_code))

                coupon_input = wait.until(EC.presence_of_element_located((By.CSS_SELECTOR, COUPON_INPUT_SELECTOR)))
                coupon_input.clear()
                coupon_input.send_keys(coupon_code)

                submit_btn = wait.until(EC.element_to_be_clickable((By.CSS_SELECTOR, SUBMIT_BUTTON_SELECTOR)))
                submit_btn.click()
                
                # alert이 아님
                # try:
                #     alert = WebDriverWait(driver, 5).until(EC.alert_is_present())
                #     msg = alert.text
                #     alert.accept()

                #     if "쿠폰함" in msg or "완료" in msg:
                #         ws[f"C{idx}"] = "성공"
                #         log_append(f"✅ {cs_code} → 성공")
                #     else:
                #         ws[f"C{idx}"] = f"실패: {msg}"
                #         log_append(f"⚠️ {cs_code} → 실패: {msg}")

                # except:
                #     ws[f"C{idx}"] = "오류: 팝업 없음/확인 필요"
                #     log_append(f"⚠️ {cs_code} → 팝업 없음 또는 실패")
                
                try:
                    # 모달이 나타날 때까지 최대 5초 대기
                    popup = WebDriverWait(driver, 5).until(
                        EC.visibility_of_element_located((By.CSS_SELECTOR, "div.pop_wrap.coupon_lyr"))
                    )
                    
                    # 팝업 메시지 가져오기
                    msg_elem = popup.find_element(By.ID, "layer_msg")
                    msg = msg_elem.text

                    # 결과 기록
                    if "완료" in msg or "쿠폰함" in msg:
                        ws[f"C{idx}"] = f"성공: {msg}"
                        log_append(f"✅ {cs_code} → 성공: {msg}")
                    else:
                        ws[f"C{idx}"] = f"실패: {msg}"
                        log_append(f"⚠️ {cs_code} → 실패: {msg}")

                    # 팝업 닫기
                    close_btn = popup.find_element(By.ID, "layer_close_btn")
                    close_btn.click()

                except:
                    ws[f"C{idx}"] = "오류: 팝업 없음/확인 필요"
                    log_append(f"⚠️ {cs_code} → 팝업 없음 또는 실패")

                time.sleep(SLEEP_SEC)

            except Exception as e:
                ws[f"C{idx}"] = f"오류: {str(e)}"
                log_append(f"❌ {cs_code} 처리 중 오류: {e}")

        result_path = excel_path.replace(".xlsx", "_result.xlsx")
        wb.save(result_path)
        driver.quit()

        lbl_status.config(text="✅ 모든 작업 완료")
        log_append("\n✅ 모든 ID 처리 완료!")
        log_append(f"📁 결과 파일 저장됨: {os.path.basename(result_path)}")
        messagebox.showinfo("완료", "✅ 모든 ID 처리 완료!\n엑셀에 결과가 기록되었습니다.")

    except Exception as e:
        log_append(f"❌ 오류 발생: {e}")
        messagebox.showerror("오류 발생", str(e))

def start_process():
    excel_path = entry_excel.get()
    coupon_code = entry_coupon.get().strip()

    if not excel_path:
        messagebox.showwarning("경고", "엑셀 파일을 선택하세요.")
        return
    if not coupon_code:
        messagebox.showwarning("경고", "쿠폰 코드를 입력하세요.")
        return

    threading.Thread(target=run_coupon_process, args=(excel_path, coupon_code), daemon=True).start()

def browse_excel():
    path = filedialog.askopenfilename(filetypes=[("Excel files", "*.xlsx")])
    if path:
        entry_excel.delete(0, tk.END)
        entry_excel.insert(0, path)

# ================= Tkinter GUI =================
root = tk.Tk()
root.title("SoulStrike 쿠폰 자동등록기")
root.geometry("600x450")
root.resizable(False, False)

frame = tk.Frame(root, padx=20, pady=20)
frame.pack(fill="both", expand=True)

tk.Label(frame, text="엑셀 파일 경로:").grid(row=0, column=0, sticky="w")
entry_excel = tk.Entry(frame, width=45)
entry_excel.grid(row=1, column=0, pady=5, sticky="w")
btn_browse = tk.Button(frame, text="찾아보기", command=browse_excel)
btn_browse.grid(row=1, column=1, padx=5)

tk.Label(frame, text="쿠폰 코드:").grid(row=2, column=0, pady=(15, 0), sticky="w")
entry_coupon = tk.Entry(frame, width=35)
entry_coupon.grid(row=3, column=0, pady=5, sticky="w")

btn_run = tk.Button(frame, text="실행하기", command=start_process, bg="#3a7afe", fg="white", width=15)
btn_run.grid(row=3, column=1, padx=5)

lbl_status = tk.Label(frame, text="대기 중", fg="gray")
lbl_status.grid(row=4, column=0, pady=(15, 5), sticky="w")

# 로그 창
tk.Label(frame, text="실행 로그:").grid(row=5, column=0, sticky="w", pady=(10, 0))
txt_log = tk.Text(frame, height=10, width=70, wrap="word", bg="#f8f8f8")
txt_log.grid(row=6, column=0, columnspan=2, pady=5)
scrollbar = tk.Scrollbar(frame, command=txt_log.yview)
scrollbar.grid(row=6, column=2, sticky="ns")
txt_log.config(yscrollcommand=scrollbar.set)

tk.Label(frame, text="💡 엑셀의 A열에 CS 코드 목록을 입력해주세요.").grid(row=7, column=0, pady=(10, 0), sticky="w")

root.mainloop()
